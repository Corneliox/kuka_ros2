import docx
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
import os
import copy

SOURCE_DOC = r"D:\Download Move1\Jurnal_lolo\From Gilang\20260901_Rev[1].docx"
OUT_REVISE = r"D:\Download Move1\Jurnal_lolo\From Gilang\Revise\Penjelasan_Rumus_Homografi.docx"
OUT_CONF = r"D:\Download Move1\Jurnal_lolo\From Gilang\conf\Penjelasan_Rumus_Homografi.docx"

def extract_and_build_standalone_word():
    src_doc = docx.Document(SOURCE_DOC)
    
    target_idx = None
    for i, p in enumerate(src_doc.paragraphs):
        if "planar homography transformation" in p.text.lower():
            target_idx = i
            break
            
    print(f"Found homography section at paragraph {target_idx}")
    
    # Create new document
    new_doc = docx.Document()
    
    # 1. Paragraph 1: Intro to Eq 1
    p1_src = src_doc.paragraphs[target_idx]
    p1_elem = copy.deepcopy(p1_src._p)
    for t in p1_elem.xpath('.//w:t'):
        if 'matric' in t.text:
            t.text = t.text.replace('matric', 'matrix')
    new_doc._body._body.append(p1_elem)
    
    # 2. Equation (1)
    eq1_src = src_doc.paragraphs[target_idx + 1]
    eq1_elem = copy.deepcopy(eq1_src._p)
    new_doc._body._body.append(eq1_elem)
    
    # 3. Paragraph 2: Connecting text between Eq 1 and Eq 2
    p2_src = src_doc.paragraphs[target_idx + 2]
    p2_elem = copy.deepcopy(p2_src._p)
    new_doc._body._body.append(p2_elem)
    
    # 4. Equation (2)
    eq2_src = src_doc.paragraphs[target_idx + 3]
    eq2_elem = copy.deepcopy(eq2_src._p)
    new_doc._body._body.append(eq2_elem)
    
    # 5. Explanatory paragraph directly UNDERNEATH Equation (2)
    temp_doc = docx.Document()
    p3 = temp_doc.add_paragraph()
    p3.paragraph_format.line_spacing = 1.15
    p3.paragraph_format.space_after = Pt(6)
    p3.paragraph_format.space_before = Pt(6)
    
    p3.add_run("Here, ")
    run_uv = p3.add_run("(u, v)")
    run_uv.italic = True
    p3.add_run(" represent the horizontal and vertical pixel coordinates of the detected target centroid on the camera image plane, whereas ")
    run_xy = p3.add_run("(Xw, Yw)")
    run_xy.italic = True
    p3.add_run(" denote the resulting physical Cartesian coordinates in millimeters relative to the robot base coordinate frame. The coefficients ")
    run_h = p3.add_run("h11 through h33")
    run_h.italic = True
    p3.add_run(" account for perspective rectification, orientation alignment, and spatial scaling between the optical sensor and the workspace table. These computed metric coordinates ")
    run_xy2 = p3.add_run("(Xw, Yw)")
    run_xy2.italic = True
    p3.add_run(" are directly transmitted to the MoveIt 2 motion planning pipeline to execute deterministic pick-and-place trajectories.")
    
    for r in p3.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    p3_elem = copy.deepcopy(p3._p)
    new_doc._body._body.append(p3_elem)
    
    # If there is a leading empty paragraph in new_doc, remove it
    if len(new_doc.paragraphs) > 5 and new_doc.paragraphs[0].text == "":
        p_first = new_doc.paragraphs[0]._p
        p_first.getparent().remove(p_first)
        
    for p in [OUT_REVISE, OUT_CONF]:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        new_doc.save(p)
        print(f"Saved standalone docx to: {p}")

if __name__ == "__main__":
    extract_and_build_standalone_word()
